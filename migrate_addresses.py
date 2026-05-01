"""
Migrate address data:
- Read Addresses sheet from customers.xlsx
- Read lat/lng from Customers sheet and create default address records for those customers
- Write everything to a new addresses.xlsx
- Clear lat/lng columns from customers.xlsx
"""
import uuid
from datetime import datetime
import openpyxl

CUSTOMERS_FILE = "data/customers.xlsx"
ADDRESSES_FILE = "data/addresses.xlsx"

def row_to_dict(ws, row_idx, headers):
    return {headers[i]: ws.cell(row=row_idx, column=i+1).value for i in range(len(headers))}

def get_headers(ws):
    return [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]

wb_c = openpyxl.load_workbook(CUSTOMERS_FILE)

# --- Read Customers sheet ---
ws_cust = wb_c["Customers"]
cust_headers = get_headers(ws_cust)
customers = []
for r in range(2, ws_cust.max_row + 1):
    row = row_to_dict(ws_cust, r, cust_headers)
    if row.get("phone"):
        customers.append(row)

# --- Read existing Addresses sheet ---
ws_addr = wb_c["Addresses"] if "Addresses" in wb_c.sheetnames else None
addr_headers = get_headers(ws_addr) if ws_addr else []
existing_addresses = []
if ws_addr:
    for r in range(2, ws_addr.max_row + 1):
        row = row_to_dict(ws_addr, r, addr_headers)
        if row.get("id") and row.get("customerId"):
            existing_addresses.append(row)

# Find customers that already have a default address in the Addresses sheet
customers_with_default = {a["customerId"] for a in existing_addresses if a.get("isDefault")}

# For each customer with lat/lng but no default address, create one
new_addresses = []
for c in customers:
    lat = c.get("latitude")
    lng = c.get("longitude")
    cid = c.get("phone")  # customerId = phone
    if lat and lng and cid not in customers_with_default:
        new_addresses.append({
            "id": str(uuid.uuid4()),
            "customerId": cid,
            "label": "Primary",
            "address": c.get("address", ""),
            "zone": c.get("zone", ""),
            "isDefault": True,
            "latitude": lat,
            "longitude": lng,
            "createdAt": c.get("createdAt") or datetime.now().isoformat(),
        })
        print(f"  Creating primary address for {c.get('name')} ({cid}): {lat}, {lng}")
    elif lat and lng and cid in customers_with_default:
        # Customer already has a default address — update its coords if it has none
        for a in existing_addresses:
            if a["customerId"] == cid and a.get("isDefault"):
                if not a.get("latitude") and not a.get("longitude"):
                    a["latitude"] = lat
                    a["longitude"] = lng
                    print(f"  Updated coords for existing default address of {c.get('name')}")
                break

all_addresses = existing_addresses + new_addresses

# --- Write addresses.xlsx ---
wb_new = openpyxl.Workbook()
ws_new = wb_new.active
ws_new.title = "Addresses"

out_headers = ["id", "customerId", "label", "address", "zone", "isDefault", "latitude", "longitude", "createdAt"]
ws_new.append(out_headers)
for a in all_addresses:
    ws_new.append([a.get(h) for h in out_headers])

wb_new.save(ADDRESSES_FILE)
print(f"\nWrote {len(all_addresses)} addresses to {ADDRESSES_FILE}")

# --- Clear lat/lng from customers.xlsx ---
lat_col = cust_headers.index("latitude") + 1 if "latitude" in cust_headers else None
lng_col = cust_headers.index("longitude") + 1 if "longitude" in cust_headers else None
if lat_col and lng_col:
    for r in range(2, ws_cust.max_row + 1):
        ws_cust.cell(row=r, column=lat_col).value = None
        ws_cust.cell(row=r, column=lng_col).value = None
    print("Cleared lat/lng from customers.xlsx")

wb_c.save(CUSTOMERS_FILE)
print("Done.")
